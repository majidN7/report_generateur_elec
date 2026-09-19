"""Import du fichier Excel de fusion des bureaux de vote.

Deux formats sont pris en charge :

1. **Format historique** (une seule feuille, ex. "Donnees_Fusion") avec les
   colonnes : الرئيس, نائب الرئيس, رقم مكتب التصويت, الجماعة, عنوان مكتب
   التصويت, العضو الأول/الثاني/الثالث, نائب العضو الأول/الثاني/الثالث —
   avec, en option, les colonnes CIN ci-dessous.

2. **Format 2026** ("Base_Fusion_Bureaux_Vote__BV.xlsx"), détecté par la
   présence d'une feuille "رؤساء وأعضاء مكاتب التصويت" et/ou "مكاتب التصويت
   المركزية" : les bureaux de vote et les bureaux centraux sont fournis dans
   deux feuilles séparées du même classeur, chacune avec le CIN de chaque
   personne. Les deux feuilles sont importées en une seule fois par ce
   module (la feuille des bureaux centraux est traitée par
   `excel_import_central.import_central_worksheet`).

L'import des bureaux de vote (les deux formats) ne lit plus رقم المكتب
المركزي / رئيس المكتب المركزي : le rattachement d'un bureau à son bureau
central se saisit désormais uniquement à la main (`numero_bureau_central`/
`president_bureau_central` restent nullable sur `BureauVote`, requis pour
générer l'arrêté — voir word_merge.py). Les bureaux centraux eux-mêmes sont
importés séparément (feuille dédiée du format 2026, ou fichier dédié via
`excel_import_central.import_excel_bureaux_centraux`).

Colonnes CIN (optionnelles au stockage, requises pour la génération) :
    بطاقة التعريف الوطنية الرئيس, بطاقة التعريف الوطنية نائب الرئيس,
    بطاقة التعريف الوطنية العضو الأول/الثاني/الثالث,
    بطاقة التعريف الوطنية نائب العضو الأول/الثاني/الثالث
"""

import io

import openpyxl
from sqlalchemy.orm import Session

from app.models.bureau_vote import BureauVote
from app.schemas.import_report import ImportReport, ImportRowError
from app.services.excel_import_central import find_central_sheet, import_central_worksheet

CORE_COLUMN_MAP = {
    "الرئيس": "president",
    "نائب الرئيس": "vice_president",
    "رقم مكتب التصويت": "numero_bureau",
    "الجماعة": "commune",
    "عنوان مكتب التصويت": "adresse_bureau",
    "العضو الأول": "membre_1",
    "العضو الثاني": "membre_2",
    "العضو الثالث": "membre_3",
    "نائب العضو الأول": "suppleant_1",
    "نائب العضو الثاني": "suppleant_2",
    "نائب العضو الثالث": "suppleant_3",
}

CIN_COLUMN_MAP = {
    "بطاقة التعريف الوطنية الرئيس": "president_cin",
    "بطاقة التعريف الوطنية نائب الرئيس": "vice_president_cin",
    "بطاقة التعريف الوطنية العضو الأول": "membre_1_cin",
    "بطاقة التعريف الوطنية العضو الثاني": "membre_2_cin",
    "بطاقة التعريف الوطنية العضو الثالث": "membre_3_cin",
    "بطاقة التعريف الوطنية نائب العضو الأول": "suppleant_1_cin",
    "بطاقة التعريف الوطنية نائب العضو الثاني": "suppleant_2_cin",
    "بطاقة التعريف الوطنية نائب العضو الثالث": "suppleant_3_cin",
}

COLUMN_MAP = {**CORE_COLUMN_MAP, **CIN_COLUMN_MAP}
REQUIRED_FIELDS = list(CORE_COLUMN_MAP.values())
OPTIONAL_FIELDS = list(CIN_COLUMN_MAP.values())

PREFERRED_SHEET_NAMES = ["رؤساء وأعضاء مكاتب التصويت", "Donnees_Fusion"]
DUAL_FORMAT_SHEET_NAMES = {"رؤساء وأعضاء مكاتب التصويت", "مكاتب التصويت المركزية"}


def _find_vote_sheet(workbook: openpyxl.Workbook):
    stripped = {name.strip(): name for name in workbook.sheetnames}
    for preferred in PREFERRED_SHEET_NAMES:
        if preferred in stripped:
            return workbook[stripped[preferred]]
    for name in workbook.sheetnames:
        ws = workbook[name]
        header_row = [c.value for c in ws[1]]
        if any(h in COLUMN_MAP for h in header_row):
            return ws
    return workbook[workbook.sheetnames[0]]


def _is_dual_sheet_format(workbook: openpyxl.Workbook) -> bool:
    stripped = {name.strip() for name in workbook.sheetnames}
    return bool(stripped & DUAL_FORMAT_SHEET_NAMES)


def _clean(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _merge_reports(reports: list[ImportReport]) -> ImportReport:
    return ImportReport(
        total_rows=sum(r.total_rows for r in reports),
        created=sum(r.created for r in reports),
        updated=sum(r.updated for r in reports),
        skipped_duplicates=sum(r.skipped_duplicates for r in reports),
        errors=[e for r in reports for e in r.errors],
        bureaux_centraux_created=sum(r.bureaux_centraux_created for r in reports),
    )


def _import_vote_worksheet(db: Session, ws) -> ImportReport:
    header_row = [c.value for c in ws[1]]
    header_index = {name: idx for idx, name in enumerate(header_row) if name in COLUMN_MAP}

    missing_columns = [h for h in CORE_COLUMN_MAP if h not in header_index]
    if missing_columns:
        raise ValueError(
            "Colonnes manquantes dans le fichier Excel (bureaux de vote) : " + ", ".join(missing_columns)
        )

    created = 0
    updated = 0
    skipped_duplicates = 0
    errors: list[ImportRowError] = []
    seen_keys: set[tuple[str, str]] = set()

    total_rows = 0
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        values = [c.value for c in row]
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        total_rows += 1

        record: dict[str, str | None] = {}
        for header, field in COLUMN_MAP.items():
            col = header_index.get(header)
            record[field] = _clean(values[col]) if col is not None and col < len(values) else None

        row_errors = [field for field in REQUIRED_FIELDS if not record.get(field)]
        if row_errors:
            errors.append(
                ImportRowError(
                    row=row_idx,
                    message=f"Champs obligatoires manquants : {', '.join(row_errors)}",
                )
            )
            continue

        key = (record["commune"], record["numero_bureau"])
        if key in seen_keys:
            skipped_duplicates += 1
            continue
        seen_keys.add(key)

        existing = (
            db.query(BureauVote)
            .filter(BureauVote.commune == key[0], BureauVote.numero_bureau == key[1])
            .one_or_none()
        )
        if existing:
            for field in REQUIRED_FIELDS:
                setattr(existing, field, record[field])
            for field in OPTIONAL_FIELDS:
                if record.get(field):
                    setattr(existing, field, record[field])
            updated += 1
        else:
            db.add(BureauVote(**record))
            created += 1

    db.commit()

    return ImportReport(
        total_rows=total_rows,
        created=created,
        updated=updated,
        skipped_duplicates=skipped_duplicates,
        errors=errors,
        bureaux_centraux_created=0,
    )


def import_excel_file(db: Session, file_bytes: bytes) -> ImportReport:
    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    if _is_dual_sheet_format(workbook):
        reports = []
        stripped = {name.strip(): name for name in workbook.sheetnames}
        if "مكاتب التصويت المركزية" in stripped:
            central_ws = find_central_sheet(workbook)
            reports.append(import_central_worksheet(db, central_ws))
        if "رؤساء وأعضاء مكاتب التصويت" in stripped:
            vote_ws = workbook[stripped["رؤساء وأعضاء مكاتب التصويت"]]
            reports.append(_import_vote_worksheet(db, vote_ws))
        return _merge_reports(reports)

    ws = _find_vote_sheet(workbook)
    return _import_vote_worksheet(db, ws)
