import shutil
import zipfile
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn

from tests.test_bureaux_crud import SAMPLE

SOFFICE_AVAILABLE = shutil.which("soffice") is not None

CIN_FIELDS = {
    "president_cin": "AB1",
    "vice_president_cin": "AB2",
    "membre_1_cin": "AB3",
    "membre_2_cin": "AB4",
    "membre_3_cin": "AB5",
    "suppleant_1_cin": "AB6",
    "suppleant_2_cin": "AB7",
    "suppleant_3_cin": "AB8",
}
SAMPLE_WITH_CIN = dict(SAMPLE, **CIN_FIELDS)

CENTRAL_CIN_FIELDS = {
    "president_cin": "PC1",
    "vice_president_cin": "PC2",
    "membre_central_1_cin": "PC3",
    "membre_central_2_cin": "PC4",
    "membre_central_3_cin": "PC5",
    "suppleant_central_1_cin": "PC6",
    "suppleant_central_2_cin": "PC7",
    "suppleant_central_3_cin": "PC8",
}


def test_generate_bureau_docx_requires_cin(client):
    bureau_id = client.post("/api/bureaux", json=SAMPLE).json()["id"]
    incomplete = client.get(f"/api/bureaux/{bureau_id}/document?format=docx")
    assert incomplete.status_code == 422
    assert "CIN" in incomplete.json()["detail"]


def test_generate_bureau_docx(client):
    bureau_id = client.post("/api/bureaux", json=SAMPLE_WITH_CIN).json()["id"]
    response = client.get(f"/api/bureaux/{bureau_id}/document?format=docx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert len(response.content) > 1000


def test_generated_docx_keeps_decision_header_static(client):
    """قرار عاملي رقم .........../2026 must appear byte-for-byte unchanged
    in the generated document, even when a numero_decision was (uselessly)
    sent in the create payload -- it's ignored by the schema and there is
    no template field for it."""
    payload = dict(SAMPLE_WITH_CIN, numero_decision="9999")
    bureau_id = client.post("/api/bureaux", json=payload).json()["id"]

    response = client.get(f"/api/bureaux/{bureau_id}/document?format=docx")
    assert response.status_code == 200

    doc = Document(BytesIO(response.content))
    header_paragraph = doc.paragraphs[0].text
    assert header_paragraph == "قرار عاملي رقم .........../2026"
    assert "9999" not in header_paragraph


@pytest.mark.skipif(not SOFFICE_AVAILABLE, reason="LibreOffice n'est pas installé")
def test_generate_bureau_pdf(client):
    bureau_id = client.post("/api/bureaux", json=SAMPLE_WITH_CIN).json()["id"]
    response = client.get(f"/api/bureaux/{bureau_id}/document?format=pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content[:4] == b"%PDF"


def test_generate_batch_zip_skips_incomplete_cin(client):
    for i in range(1, 4):
        client.post("/api/bureaux", json=dict(SAMPLE_WITH_CIN, numero_bureau=str(i)))
    # one bureau without CIN, should be skipped rather than fail the whole batch
    client.post("/api/bureaux", json=dict(SAMPLE, numero_bureau="99"))

    response = client.post("/api/bureaux/generate-batch?format=docx")
    assert response.status_code == 200
    assert response.headers["X-Skipped-Incomplete"] == "1"
    zf = zipfile.ZipFile(BytesIO(response.content))
    assert len(zf.namelist()) == 3


def test_generated_documents_keep_the_official_seal_image(client):
    """Regression test: the official seal/stamp is embedded as a <w:drawing>
    anchored on the signature line. On bureau_central.docx that drawing
    shares its run with the "الداخلة، في: ..." text, so a naive rebuild of
    that paragraph's runs (to inject date_signature) would silently delete
    it. Both generated documents must keep at least one drawing."""
    vote_id = client.post("/api/bureaux", json=SAMPLE_WITH_CIN).json()["id"]
    vote_doc = client.get(f"/api/bureaux/{vote_id}/document?format=docx")
    assert vote_doc.status_code == 200
    vote_drawings = Document(BytesIO(vote_doc.content)).element.body.findall(".//" + qn("w:drawing"))
    assert len(vote_drawings) >= 1

    central = client.post(
        "/api/bureaux-centraux",
        json={
            "numero_bureau_central": "1",
            "commune": "الداخلة",
            "president_bureau_central": "عبد الحق بلعابد",
            "adresse_bureau_central": "مقر المكتب المركزي",
            "vice_president_bureau_central": "نائب الرئيس",
            "membre_central_1": "عضو 1",
            "membre_central_2": "عضو 2",
            "membre_central_3": "عضو 3",
            "suppleant_central_1": "نائب 1",
            "suppleant_central_2": "نائب 2",
            "suppleant_central_3": "نائب 3",
            **CENTRAL_CIN_FIELDS,
        },
    ).json()
    central_doc = client.get(f"/api/bureaux-centraux/{central['id']}/document?format=docx")
    assert central_doc.status_code == 200
    central_drawings = Document(BytesIO(central_doc.content)).element.body.findall(".//" + qn("w:drawing"))
    assert len(central_drawings) >= 1


def test_bureau_central_document_requires_complete_data(client):
    central = client.post(
        "/api/bureaux-centraux",
        json={
            "numero_bureau_central": "1",
            "commune": "الداخلة",
            "president_bureau_central": "عبد الحق بلعابد",
        },
    ).json()

    incomplete = client.get(f"/api/bureaux-centraux/{central['id']}/document?format=docx")
    assert incomplete.status_code == 422

    complete_payload = {
        "adresse_bureau_central": "مقر المكتب المركزي",
        "vice_president_bureau_central": "نائب الرئيس",
        "membre_central_1": "عضو 1",
        "membre_central_2": "عضو 2",
        "membre_central_3": "عضو 3",
        "suppleant_central_1": "نائب 1",
        "suppleant_central_2": "نائب 2",
        "suppleant_central_3": "نائب 3",
        **CENTRAL_CIN_FIELDS,
    }
    client.put(f"/api/bureaux-centraux/{central['id']}", json=complete_payload)

    complete = client.get(f"/api/bureaux-centraux/{central['id']}/document?format=docx")
    assert complete.status_code == 200
    assert len(complete.content) > 1000
